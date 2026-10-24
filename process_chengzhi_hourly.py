from datetime import datetime
from pathlib import Path
import json
import subprocess
import urllib.request

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


BASE_DIR = Path(__file__).resolve().parent
SOURCE_CSV = BASE_DIR / "余杭区良渚古城遗址公园综合管控系统-游客出入园信息20260618.csv"
OUTPUT_CSV = BASE_DIR / "城址区_按小时汇总_实际入园游客量.csv"
OUTPUT_XLSX = BASE_DIR / "城址区_按小时汇总_实际入园游客量.xlsx"
HOLIDAY_CACHE = BASE_DIR / "china_holiday_2019_2026.json"
WEATHER_CACHE = BASE_DIR / "hangzhou_hourly_weather_2019-06-26_2026-06-02.json"

HANGZHOU_LATITUDE = 30.2741
HANGZHOU_LONGITUDE = 120.1551

WEEKDAY_NAMES = {
    0: "星期一",
    1: "星期二",
    2: "星期三",
    3: "星期四",
    4: "星期五",
    5: "星期六",
    6: "星期日",
}

WEATHER_CODE_NAMES = {
    0: "晴",
    1: "大部晴朗",
    2: "局部多云",
    3: "阴",
    45: "雾",
    48: "雾凇",
    51: "小毛毛雨",
    53: "中毛毛雨",
    55: "大毛毛雨",
    56: "冻毛毛雨",
    57: "强冻毛毛雨",
    61: "小雨",
    63: "中雨",
    65: "大雨",
    66: "冻雨",
    67: "强冻雨",
    71: "小雪",
    73: "中雪",
    75: "大雪",
    77: "雪粒",
    80: "小阵雨",
    81: "中阵雨",
    82: "强阵雨",
    85: "小阵雪",
    86: "强阵雪",
    95: "雷暴",
    96: "雷暴伴小冰雹",
    99: "雷暴伴大冰雹",
}


def download_json(url: str) -> dict:
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            return json.loads(response.read().decode("utf-8"))
    except Exception:
        command = (
            "$ProgressPreference='SilentlyContinue'; "
            f"Invoke-RestMethod -Uri '{url}' | ConvertTo-Json -Depth 20"
        )
        result = subprocess.run(
            ["powershell", "-NoProfile", "-Command", command],
            check=True,
            capture_output=True,
            text=True,
            encoding="utf-8",
        )
        return json.loads(result.stdout)


def load_holiday_flags(start_year: int, end_year: int) -> dict:
    if HOLIDAY_CACHE.exists():
        return json.loads(HOLIDAY_CACHE.read_text(encoding="utf-8"))

    holiday_flags = {}
    for year in range(start_year, end_year + 1):
        url = f"https://cdn.jsdelivr.net/gh/NateScarlet/holiday-cn@master/{year}.json"
        data = download_json(url)
        for day in data.get("days", []):
            holiday_flags[day["date"]] = bool(day["isOffDay"])

    HOLIDAY_CACHE.write_text(
        json.dumps(holiday_flags, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return holiday_flags


def load_weather_by_hour(start_date: str, end_date: str) -> dict:
    if WEATHER_CACHE.exists():
        return json.loads(WEATHER_CACHE.read_text(encoding="utf-8"))

    url = (
        "https://archive-api.open-meteo.com/v1/archive"
        f"?latitude={HANGZHOU_LATITUDE}"
        f"&longitude={HANGZHOU_LONGITUDE}"
        f"&start_date={start_date}"
        f"&end_date={end_date}"
        "&hourly=weather_code"
        "&timezone=Asia%2FShanghai"
    )
    data = download_json(url)
    times = data["hourly"]["time"]
    codes = data["hourly"]["weather_code"]
    weather_by_hour = {
        pd.to_datetime(time).strftime("%Y-%m-%d %H:%M:%S"): WEATHER_CODE_NAMES.get(code, f"天气代码{code}")
        for time, code in zip(times, codes)
    }

    WEATHER_CACHE.write_text(
        json.dumps(weather_by_hour, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return weather_by_hour


def build_hourly_summary() -> pd.DataFrame:
    df = pd.read_csv(SOURCE_CSV, encoding="utf-8-sig")

    filtered = df.loc[df["区域"].astype(str).str.strip() == "0", ["统计时间", "实际入园游客量"]].copy()
    filtered["统计时间"] = pd.to_datetime(filtered["统计时间"]).dt.floor("h")
    filtered["实际入园游客量"] = pd.to_numeric(filtered["实际入园游客量"], errors="coerce").fillna(0)

    summary = (
        filtered.groupby("统计时间", as_index=False)["实际入园游客量"]
        .sum()
        .sort_values("统计时间")
        .rename(columns={"统计时间": "create_time", "实际入园游客量": "entrance_num"})
    )
    summary["entrance_num"] = summary["entrance_num"].astype(int)
    start_date = summary["create_time"].min().strftime("%Y-%m-%d")
    end_date = summary["create_time"].max().strftime("%Y-%m-%d")
    holiday_flags = load_holiday_flags(
        summary["create_time"].dt.year.min(),
        summary["create_time"].dt.year.max(),
    )
    weather_by_hour = load_weather_by_hour(start_date, end_date)

    summary["星期"] = summary["create_time"].dt.weekday.map(WEEKDAY_NAMES)
    summary["是否节假日"] = summary["create_time"].dt.strftime("%Y-%m-%d").map(holiday_flags).fillna(False)
    summary["是否节假日"] = summary["是否节假日"].map({True: "是", False: "否"})
    summary["天气"] = summary["create_time"].dt.strftime("%Y-%m-%d %H:%M:%S").map(weather_by_hour).fillna("未知")
    return summary


def save_outputs(summary: pd.DataFrame) -> Path:
    summary.to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")

    xlsx_path = OUTPUT_XLSX
    try:
        summary.to_excel(xlsx_path, index=False, sheet_name="城址区按小时汇总")
    except PermissionError:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        xlsx_path = OUTPUT_XLSX.with_name(f"{OUTPUT_XLSX.stem}_{timestamp}{OUTPUT_XLSX.suffix}")
        summary.to_excel(xlsx_path, index=False, sheet_name="城址区按小时汇总")

    wb = load_workbook(xlsx_path)
    ws = wb["城址区按小时汇总"]
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 14

    header_fill = PatternFill("solid", fgColor="1F6F78")
    header_font = Font(bold=True, color="FFFFFF")
    thin_border = Side(style="thin", color="D9E2E3")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = Border(bottom=thin_border)

    for cell in ws["A"][1:]:
        cell.number_format = "yyyy-mm-dd hh:mm:ss"
    for cell in ws["B"][1:]:
        cell.number_format = "#,##0"

    table_ref = f"A1:E{ws.max_row}"
    table = Table(displayName="HourlyEntranceSummary", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    wb.save(xlsx_path)
    return xlsx_path


def main() -> None:
    summary = build_hourly_summary()
    xlsx_path = save_outputs(summary)

    print(f"source_csv: {SOURCE_CSV}")
    print(f"output_csv: {OUTPUT_CSV}")
    print(f"output_xlsx: {xlsx_path}")
    print(f"output_hours: {len(summary)}")
    print(f"duplicate_hours: {summary['create_time'].duplicated().sum()}")
    print(f"total_entrance_num: {int(summary['entrance_num'].sum())}")
    print(f"unknown_weather_rows: {(summary['天气'] == '未知').sum()}")


if __name__ == "__main__":
    main()
